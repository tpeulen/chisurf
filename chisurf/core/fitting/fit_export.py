"""Write a fit to a folder, and read one back: *File ▸ Fits* (Load Fit, Current Fit, All Fits), without Qt.

What a fit export writes beside each other, named after the fit's data (``<stem>``):

* the curves of every member (``<stem>_<member>_<curve>.csv``, ``…_info.txt``): :meth:`FitGroup.save`;
* ``<stem>.fit.cs.pto``: the fit and its data as a ChiSurf document (the ``.cs.pto`` container a project is
  written in, marked ``document_kind: fit``). *Load Fit* adds it to an open session without replacing anything;
* ``<stem>.docx``: a report -- the pictures it is given (each fit's plots and model, drawn by emtk:
  :func:`chisurf.emtk.fit_report.render_fit_images`) and a table of every parameter of every member, fitted
  ones in bold, linked ones in italics, with χ²ᵣ.

*All Fits* writes one folder per fit (global fits are left out: their members are written as fits of their own)
and stacks the reports into ``all_fits.docx``.

Nothing here draws or imports a GUI toolkit. The pictures come in as PNG bytes.
"""

from __future__ import annotations

import logging
import os
import shutil
from collections.abc import Callable, Iterable, Sequence
from pathlib import Path
from typing import Any, TypedDict

_log = logging.getLogger("chisurf.fitting.fit_export")

#: The suffix of a fit document (a ``.cs.pto`` holding one fit and its data).
FIT_DOCUMENT_SUFFIX = ".fit.cs.pto"
#: The stacked report *All Fits* writes in the target folder.
ALL_FITS_REPORT = "all_fits.docx"

Images = Sequence[tuple[str, bytes]]


def members(fit: Any) -> list[Any]:
    """The local fits of *fit* (a fit group's members, or the fit itself)."""
    grouped = getattr(fit, "grouped_fits", None)
    return list(grouped) if grouped else [fit]


def is_global_fit(fit: Any) -> bool:
    """Whether *fit* is a global fit (its data is the global dataset's setup)."""
    from chisurf.core.experiments.globalfit import GlobalFitSetup

    return isinstance(getattr(getattr(fit, "data", None), "setup", None), GlobalFitSetup)


def fit_stem(fit: Any, use_complex_name: bool = False) -> str:
    """The base name of a fit's files: its data's file name without extension (or the fit's own name)."""
    from chisurf.core.base import clean_string

    if use_complex_name:
        name = clean_string(str(getattr(fit, "name", "") or "fit"))
    else:
        name = os.path.basename(str(getattr(getattr(fit, "data", None), "name", "") or "fit"))
    return os.path.splitext(name)[0] or "fit"


def write_fit_document(fit: Any, path: str | Path) -> Path:
    """Write *fit* and its data as a fit document (``.cs.pto``) at *path*."""
    from chisurf.core.data import DataCurveGroup
    from chisurf.core.project import capture_session
    from chisurf.core.project.storage import save_file

    name = str(getattr(fit, "name", "") or Path(path).stem)
    curves = [member.data for member in members(fit)]
    # The data keep their name: one curve is "fcs_demo", not the fit's long name, when it comes back.
    group = str(getattr(curves[0], "name", "") or name) if len(curves) == 1 else name
    project = capture_session([DataCurveGroup(curves, name=group)], [fit], name=name)
    project.metadata["document_kind"] = "fit"
    return Path(save_file(project, str(path)))


def write_report(fit: Any, path: str | Path, images: Images = ()) -> Path:
    """Write the Word report of *fit*: its *images* (caption, PNG bytes) and the parameter table."""
    import io

    import docx
    from docx.shared import Inches

    document = docx.Document()
    document.add_heading(str(getattr(fit, "name", "") or Path(path).stem), 0)
    if images:
        document.add_heading("Fit-Results", level=1)
        for caption, png in images:
            document.add_paragraph(str(caption))
            document.add_picture(io.BytesIO(png), width=Inches(5.5))

    group = members(fit)
    document.add_heading("Summary", level=1)
    paragraph = document.add_paragraph("Parameters which are fitted are given in ")
    paragraph.add_run("bold").bold = True
    paragraph.add_run(", linked parameters in ")
    paragraph.add_run("italic.").italic = True
    paragraph.add_run(" Fixed parameters are plain.")
    table = document.add_table(rows=1, cols=len(group) + 1)
    header = table.rows[0].cells
    header[0].text = "Param"
    for column in range(len(group)):
        header[column + 1].text = str(column + 1)
    names = sorted({name for member in group for name in member.model.parameters_all_dict})
    for name in names:
        row = table.add_row().cells
        row[0].text = name
        for column, member in enumerate(group):
            parameter = member.model.parameters_all_dict.get(name)
            if parameter is None:
                continue
            run = row[column + 1].paragraphs[0].add_run(f"{parameter.value:.3f}")
            if parameter.fixed:
                continue
            if parameter.link is not None:
                run.italic = True
            else:
                run.bold = True
    chi_row = table.add_row().cells
    chi_row[0].text = "Chi2r"
    for column, member in enumerate(group):
        chi_row[column + 1].paragraphs[0].add_run(f"{float(member.chi2r):.4f}")
    target = Path(path)
    document.save(str(target))
    return target


class FitExport(TypedDict):
    """What :func:`export_fit` wrote: the curve tables, the fit document, the report."""

    csv: list[str]
    document: str
    report: str


def export_fit(
    fit: Any, target_path: str | Path, images: Images = (), *, use_complex_name: bool = False
) -> FitExport:
    """*Current Fit*: write *fit*'s curves, its fit document and its report into *target_path*.

    Returns
    -------
    dict
        ``{"csv": [paths], "document": path, "report": path}``: what was written.
    """
    folder = Path(target_path)
    folder.mkdir(parents=True, exist_ok=True)
    base = folder / fit_stem(fit, use_complex_name)
    fit.save(str(base), "csv", save_curves=True)
    document = write_fit_document(fit, str(base) + FIT_DOCUMENT_SUFFIX)
    report = write_report(fit, str(base) + ".docx", images)
    _log.info("fit written to %s", folder)
    tables = sorted(str(path) for path in folder.glob(f"{base.name}*.csv"))
    return {"csv": tables, "document": str(document), "report": str(report)}


def export_fits(
    fits: Iterable[Any],
    target_path: str | Path,
    images_for: Callable[[Any], Images] | None = None,
    *,
    existing: str = "skip",
) -> dict[str, Any]:
    """*All Fits*: one folder per fit under *target_path*, and the reports stacked in ``all_fits.docx``.

    Parameters
    ----------
    fits : iterable
        The session's fits; global fits are left out.
    target_path : path
        The folder the per-fit folders go into.
    images_for : callable, optional
        ``images_for(fit) -> [(caption, png)]``, the report's pictures of one fit.
    existing : {"skip", "overwrite"}
        What happens to a fit whose folder is already there.

    Returns
    -------
    dict
        ``{"written": [folder], "skipped": [folder], "report": path or None}``.
    """
    root = Path(target_path)
    root.mkdir(parents=True, exist_ok=True)
    written, skipped, reports = [], [], []
    for fit in fits:
        if is_global_fit(fit):
            continue
        folder = root / fit_stem(fit)
        if folder.exists():
            if existing != "overwrite":
                skipped.append(str(folder))
                continue
            shutil.rmtree(folder) if folder.is_dir() else folder.unlink()
        paths = export_fit(fit, folder, images_for(fit) if images_for else ())
        written.append(str(folder))
        reports.append(paths["report"])
    report = stack_reports(reports, root / ALL_FITS_REPORT) if reports else None
    return {"written": written, "skipped": skipped, "report": str(report) if report else None}


def stack_reports(paths: Sequence[str | Path], out_path: str | Path) -> Path:
    """Stack Word reports into one document, a page break between them."""
    from copy import deepcopy

    from docx import Document

    master = Document()
    for index, path in enumerate(paths):
        if index:
            master.add_page_break()
        for child in list(Document(str(path)).element.body):
            master.element.body.append(deepcopy(child))
    target = Path(out_path)
    master.save(str(target))
    return target


def existing_folders(fits: Iterable[Any], target_path: str | Path) -> list[str]:
    """The per-fit folders *All Fits* would write that are already there (to ask before overwriting)."""
    root = Path(target_path)
    return [
        str(root / fit_stem(fit))
        for fit in fits
        if not is_global_fit(fit) and (root / fit_stem(fit)).exists()
    ]


def read_fit_document(path: str | Path):
    """*Load Fit*: the datasets and fits of a fit document (or a project file), detached.

    Returns
    -------
    chisurf.core.project.session.RestoredSession
        New objects, not yet in any session; :func:`import_fit` adds them.
    """
    from chisurf.core.project import restore_session
    from chisurf.core.project.storage import load_file

    return restore_session(load_file(str(path)))


def import_fit(path: str | Path, datasets: list, fits: list):
    """Add the fits and data of the document at *path* to *datasets* and *fits*; nothing is replaced.

    Raises
    ------
    ValueError
        The document's fits or datasets are already in the session (it was loaded before).
    """
    restored = read_fit_document(path)

    def uid(obj: Any) -> str:
        return str(getattr(obj, "unique_identifier", "") or "")

    present = {uid(obj) for obj in [*datasets, *fits]} - {""}
    incoming = {uid(obj) for obj in [*restored.datasets, *restored.fits]} - {""}
    if present & incoming:
        raise ValueError("This fit is already open in the session.")
    datasets.extend(restored.datasets)
    fits.extend(restored.fits)
    return restored


__all__ = [
    "ALL_FITS_REPORT",
    "FIT_DOCUMENT_SUFFIX",
    "existing_folders",
    "export_fit",
    "export_fits",
    "fit_stem",
    "import_fit",
    "is_global_fit",
    "members",
    "read_fit_document",
    "stack_reports",
    "write_fit_document",
    "write_report",
]
