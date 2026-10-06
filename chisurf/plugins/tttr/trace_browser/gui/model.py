"""Qt-free state and actions of the Trace Browser.

:class:`TraceBrowserModel` holds what the Qt workspace
(:class:`chisurf.plugins.tttr.trace_browser.widget.TraceBrowser`) used to keep
inside its widgets: the current folder, the include-subfolders flag, the setup
(detector / channel definition) and the channels derived from it, the file rows
(name, path, size, rating, notes), the rating filter, the bin window, the y range,
the selection and the loaded trace.  The work the Qt tool did on that state is
here as plain methods -- scan a folder and fill the rows, filter them, read and
write ratings and notes (persisted through :mod:`..core.metadata`), clear, clear
the caches, and compute or load a trace (through :mod:`..core.trace` and the
RPC client, with the in-memory and on-disk caches) -- so the Qt tool and a later
native emtk app share one implementation.

Heavy methods (:meth:`~TraceBrowserModel.scan`,
:meth:`~TraceBrowserModel.precompute_all_traces`,
:meth:`~TraceBrowserModel.load_trace`) run headless and inside
:class:`chisurf.emtk.jobs.SnapshotJob`; the model provides ``notify(event)``,
``_observers`` and ``add_observer`` for that.

The module imports no Qt binding.  Computing a trace that is not cached still
needs the binning engine of the ``intensity_trace`` plugin (a Qt widget class,
imported lazily by :func:`..core.trace.load_trace` on a cache miss only).
"""

from __future__ import annotations

import csv
import importlib
import importlib.util
import logging
import pathlib
import shutil
import tempfile
import time
from collections.abc import Callable
from typing import Any

import numpy as np

from chisurf.plugins.tttr.trace_browser.core import metadata as _metadata
from chisurf.plugins.tttr.trace_browser.core import trace as _trace

try:
    import tttrlib
except Exception:  # pragma: no cover - tttrlib is an optional dependency here
    tttrlib = None

logger = logging.getLogger(__name__)

#: Labels of the rating filter, in the order of ``filter_index`` (0..4).
FILTER_LABELS = (
    "All",
    "≥ 1★",
    "≥ 2★★",
    "≥ 3★★★",
    "Only 0★",
)

#: Highest rating (a whole number 0..this); the Qt star widget has 3 stars.
RATING_MAX = 3

#: Name of the per-folder trace cache directory (beside the data files).
CACHE_DIRNAME = ".tttr_trace_cache"

#: Container / format name of a setup -> typical file extensions.
FILETYPE_EXTENSIONS: dict[str, set[str]] = {
    "PTU": {".ptu"},
    "PT3": {".pt3"},
    "HT3": {".ht3"},
    "PT2": {".pt2"},
    "PT5": {".pt5"},
    "SPC-130": {".spc"},
    "SPC-600": {".spc"},
    "SPC-830": {".spc"},
    "PHU": {".phu"},
    "PHOTON_HDF5": {".h5", ".hdf5", ".photon.hdf5"},
    "HDF5": {".h5", ".hdf5"},
}


#: Names of the hand-off requests the model records (and a host fulfils).
HANDOFF_INTENSITY_TRACE = "open_intensity_trace"
HANDOFF_TIME_WINDOW = "open_time_window"
HANDOFF_NDXPLORER = "open_ndxplorer"

#: Spec actions that need at least one selected file.
_SELECTION_ACTIONS = frozenset(
    {
        "export_selected",
        "export_csv",
        "export_docx",
        "delete_selected",
        "open_intensity_trace",
        "open_time_window",
        "open_ndxplorer",
    }
)
_HANDOFF_ACTIONS = frozenset({"open_intensity_trace", "open_time_window", "open_ndxplorer"})


def _docx_modules() -> tuple[Any, Any] | None:
    """Return ``(Document, Inches)`` of the optional ``python-docx`` package, or ``None``."""
    try:
        from docx import Document
        from docx.shared import Inches
    except Exception:
        return None
    return Document, Inches


def get_tttr_supported_exts() -> list[str]:
    """Return the file extensions ``tttrlib`` reads (lower case, with the dot).

    Returns
    -------
    list of str
        Normalised extensions; empty when ``tttrlib`` is unavailable.
    """
    exts: list[str] = []
    try:
        if tttrlib is not None and hasattr(tttrlib, "get_supported_filetypes"):
            exts = list(tttrlib.get_supported_filetypes())
    except Exception:
        exts = []

    norm: list[str] = []
    for e in exts:
        s = str(e).strip().lower()
        if not s:
            continue
        if not s.startswith("."):
            s = "." + s
        if s not in norm:
            norm.append(s)
    return norm


class TraceBrowserModel:
    """State and actions of the Trace Browser.

    Parameters
    ----------
    client : TraceBrowserClient, optional
        RPC client for metadata and trace loading; a local in-process one is
        created on first use.
    """

    def __init__(self, client: Any = None) -> None:
        self._client = client
        self._observers: list[Callable[[Any], None]] = []
        # -- folder and setup ------------------------------------------
        #: The folder whose traces are listed (``None`` before one is opened).
        self.current_folder: pathlib.Path | None = None
        #: Also list files in sub-folders.
        self.include_subfolders = False
        #: Detector / channel definition exactly as the setup page returns it.
        self.setup_settings: dict | None = None
        #: File type of the selected setup (``None`` or ``"Auto"`` = every supported type).
        self.setup_filetype: str | None = None
        #: Routing channels used by the traces (union over the setup's detectors).
        self.selected_channels: list[int] | None = None
        #: The page shown: ``"setup"`` (detector setup) or ``"browser"``.
        self.page = "setup"
        # -- files -----------------------------------------------------
        #: Per-folder metadata ``{relative path: {"rating": int, "annotation": str}}``.
        self.meta: dict[str, dict] = {}
        #: Every scanned file as a row dict (name, path, size, size_text, size_mb, rating, notes).
        self.files: list[dict[str, Any]] = []
        #: The rows that pass the rating filter (the file table's source).
        self.rows: list[dict[str, Any]] = []
        #: Index into :data:`FILTER_LABELS`.
        self.filter_index = 0
        # -- trace view ------------------------------------------------
        #: Bin window of the trace in milliseconds.
        self.window_ms = 10.0
        self.y_min = 0.0
        self.y_max = 1000.0
        #: Paths (str) of the selected rows.
        self.selected_files: list[str] = []
        #: The file whose trace is shown.
        self.current_file: pathlib.Path | None = None
        #: The shown trace: ``{"time_axis", "counts", "labels", "time_window_ms"}``.
        self.trace: dict[str, Any] | None = None
        #: Hand-off requests for the host application: ``{"name": ..., "payload": {...}}``.
        self.requests: list[dict[str, Any]] = []
        #: ``True`` when a host (the ChiSurf main window) fulfils hand-off requests; the app sets it.
        self.host_connected = False
        #: Paths (str) of every selected row (the table's multi-row selection, set by the app);
        #: empty means "only :attr:`selected_files`".
        self.multi_selection: list[str] = []
        #: Set by :meth:`select_all_files`; the app selects every row of the table and clears it.
        self.select_all_requested = False
        #: A pending confirmation (``{"kind", "title", "message", "paths"}``) or ``None``.
        self.confirm: dict[str, Any] | None = None
        #: Number of exports, deletions and hand-offs that finished (a guided tour watches it).
        self.actions_done = 0
        self.status_text = "Open a folder with TTTR files."
        self.error_text = ""
        #: Why the trace of the current file could not be loaded (empty when it could).
        self.trace_error = ""
        #: Compute the traces of every listed file in the background after a scan (the Qt
        #: browser did this after every scan, on a modal progress dialog).
        self.precompute_after_scan = True
        #: Set by :meth:`scan` and :meth:`precompute_traces`; the host starts the job and clears it.
        self.precompute_pending = False
        #: State of the precompute job; a dict that stays the same object, so the host's worker
        #: and the draw loop see each other's writes.
        self.precompute: dict[str, Any] = {
            "running": False,
            "cancel": False,
            "done": 0,
            "total": 0,
            "name": "",
            "message": "",
        }
        #: A dialog the host app should open (``"folder"``); the app clears it when it does.
        self.dialog = ""
        #: ``callable(method, *args)`` of a host that runs heavy methods off the draw thread
        #: (the app's job queue); ``None`` runs them inline.
        self.runner: Callable[..., Any] | None = None
        #: ``True`` while the host's worker runs (the form greys its controls).
        self.busy = False
        #: ``callable(path) -> bool`` replacing the CLSM image probe (tests, hosts).
        self.image_probe: Callable[[pathlib.Path], bool] | None = None
        self._trace_mem_cache: dict[tuple[pathlib.Path, str], tuple] = {}
        self._is_image_cache: dict[pathlib.Path, bool] = {}

    # ── observers (SnapshotJob) ───────────────────────────────────────
    def add_observer(self, callback: Callable[[Any], None]) -> None:
        """Register ``callback(event)`` for :meth:`notify`."""
        self._observers.append(callback)

    def notify(self, event: Any = "updated") -> None:
        """Tell every observer that *event* happened."""
        for callback in list(self._observers):
            callback(event)

    # ── the backend client ────────────────────────────────────────────
    @property
    def client(self) -> Any:
        """The RPC client, created on first use."""
        if self._client is None:
            from .client import TraceBrowserClient

            self._client = TraceBrowserClient()
        return self._client

    # ── setup ─────────────────────────────────────────────────────────
    def apply_setup(self, settings: dict, filetype: str | None = None) -> None:
        """Accept a detector setup and derive the selected channels from it.

        Parameters
        ----------
        settings : dict
            Setup as returned by the detector wizard (``{"detectors": {...}, ...}``).
        filetype : str, optional
            The setup's file type; ``None`` / ``"Auto"`` accepts every supported type.
        """
        self.setup_settings = settings
        self.setup_filetype = filetype
        chs: list[int] = []
        for det in (settings or {}).get("detectors", {}).values():
            for c in det.get("chs", []):
                if c not in chs:
                    chs.append(c)
        self.selected_channels = sorted(chs) if chs else None
        logger.debug("TraceBrowser: Selected channels = %s", self.selected_channels)
        self.notify("setup")

    @staticmethod
    def filetype_of(settings: dict | None) -> str | None:
        """Return the file type a setup's TTTR-reading block selects.

        Parameters
        ----------
        settings : dict, optional
            Setup as returned by the setup editor.

        Returns
        -------
        str or None
            The ``tttr_reading.file_type``; ``None`` for ``Auto`` or when absent (as
            the Qt setup page's ``filetype`` property).
        """
        reading = (settings or {}).get("tttr_reading") or {}
        text = str(reading.get("file_type") or "").strip()
        return None if not text or text.lower() == "auto" else text

    def accept_setup(self, settings: dict) -> None:
        """Apply the setup page's result and open the browser page (Continue).

        Parameters
        ----------
        settings : dict
            Setup as returned by the setup editor's ``get_settings()``.
        """
        self.apply_setup(settings, self.filetype_of(settings))
        self.page = "browser"
        self.notify("page")

    def back_to_setup(self) -> None:
        """Return to the setup page; the accepted setup stays until the next Continue."""
        self.page = "setup"
        self.notify("page")

    def build_channel_labels(self, chs: list[int]) -> list[str]:
        """Return one label per routing channel in *chs*, built from the setup.

        A label is ``"<detector name>, start-end[;start2-end2]"`` (several detectors
        sharing a channel are joined with `` + ``); a channel the setup does not
        name is labelled by its number.

        Parameters
        ----------
        chs : list of int
            Routing channels.

        Returns
        -------
        list of str
            The labels, in the order of *chs*.
        """
        # Build labels like "<detector_name>, start-end[;start2-end2]" for each routing channel
        try:
            settings = self.setup_settings or {}
            dets = settings.get("detectors", {}) if isinstance(settings, dict) else {}
            # Map routing channel -> list of label parts (in case multiple detectors include same channel)
            label_map: dict[int, list[str]] = {}
            for det_name, dinfo in dets.items():
                try:
                    det_chs = list(dinfo.get("chs", []))
                    mtrs = dinfo.get("micro_time_ranges", []) or []
                    # Build range text
                    rng_txt = ";".join(
                        f"{int(a)}-{int(b)}"
                        for (a, b) in mtrs
                        if isinstance(a, (int, float)) and isinstance(b, (int, float))
                    )
                    base = det_name if det_name is not None else ""
                    lbl = f"{base}, {rng_txt}" if rng_txt else base
                    for ch in det_chs:
                        label_map.setdefault(int(ch), []).append(lbl)
                except Exception:
                    continue
            labels: list[str] = []
            for ch in chs:
                parts = label_map.get(int(ch))
                if parts:
                    # Deduplicate identical parts while preserving order
                    seen = set()
                    uniq = []
                    for p in parts:
                        if p not in seen:
                            seen.add(p)
                            uniq.append(p)
                    labels.append(" + ".join(uniq))
                else:
                    labels.append(str(ch))
            return labels
        except Exception:
            # Fallback: just stringify channels
            return [str(c) for c in chs]

    def allowed_exts(self) -> set[str]:
        """Return the file extensions allowed by the selected setup's file type.

        Returns
        -------
        set of str
            Lower-case extensions with the dot; every ``tttrlib``-supported one
            for ``Auto`` / no setup file type or an unknown one.
        """
        try:
            filetype = self.setup_filetype
            all_exts = set(get_tttr_supported_exts())
            if not filetype or str(filetype).strip().lower() == "auto":
                return all_exts
            exts = FILETYPE_EXTENSIONS.get(str(filetype).strip().upper())
            if exts:
                # Intersect with actually supported to be safe
                return {e for e in exts if (not all_exts or e in all_exts)} or exts
            return all_exts
        except Exception:
            return set(get_tttr_supported_exts())

    # ── spec-facing helpers (the emtk Browser page) ───────────────────
    def request(self, method: str, *args: Any) -> Any:
        """Run the model method *method*: through :attr:`runner` when a host set one, else now.

        A host that keeps its draw loop responsive (the emtk app) queues the call on a
        worker; headless callers and tests run it inline.
        """
        if self.runner is not None:
            return self.runner(method, *args)
        return getattr(self, method)(*args)

    @property
    def folder_text(self) -> str:
        """The folder label of the Browser page (the Qt label's text)."""
        return str(self.current_folder) if self.current_folder else "No folder selected"

    @property
    def status_line(self) -> str:
        """One line for the status area: the error, else what the last scan found."""
        if self.busy:
            return "Working..."
        text = self.error_text or self.trace_error or self.status_text
        background = self.background_text
        return f"{text}   |   {background}" if background else text

    @property
    def background_text(self) -> str:
        """The precompute job's progress while it runs, else what the last run did."""
        state = self.precompute
        if state["running"]:
            if state["total"]:
                return f"Precomputing traces {state['done'] + 1}/{state['total']}: {state['name']}"
            return "Precomputing traces..."
        return str(state["message"])

    def filter_label_list(self) -> list[str]:
        """The rating filter's drop-down entries (:data:`FILTER_LABELS`)."""
        return list(FILTER_LABELS)

    def enabled(self, name: str) -> bool:
        """Whether the control *name* may be used now.

        Every control waits for a running scan, except *Stop*, which is live only while the
        precompute job runs; *Precompute* needs listed files and no precompute running.
        """
        if name == "stop_precompute":
            return bool(self.precompute["running"])
        if self.busy or self.confirm is not None:
            return False
        if name == "precompute_traces":
            return bool(self.rows) and not self.precompute["running"]
        if name == "select_all_files":
            return bool(self.rows)
        if name in _SELECTION_ACTIONS and not self.selection_paths():
            return False
        if name == "export_docx":
            return _docx_modules() is not None
        if name in _HANDOFF_ACTIONS:
            if not self.host_connected:
                return False
            if name == HANDOFF_NDXPLORER:
                return tttrlib is not None and self.ndx_available()
        return True

    def choose_folder(self) -> None:
        """Ask the host to open its folder chooser (the Qt *Open*)."""
        self.dialog = "folder"

    def on_paths_dropped(self, paths: Any) -> bool:
        """Open the first dropped folder (the Qt ``dropEvent``); anything else is ignored.

        Returns
        -------
        bool
            ``True`` when a folder was taken.
        """
        for item in paths or ():
            try:
                folder = pathlib.Path(str(item))
                if folder.exists() and folder.is_dir():
                    self.request("open_folder", folder)
                    return True
            except Exception:
                continue
        self.status_text = "Drop a folder to open it; a file is ignored."
        return False

    def select_row(self, record: Any) -> None:
        """Select the table row *record* (the table's ``selected_call``; ``None`` clears).

        Only the selection and the current file are set; no trace is computed here.
        """
        path = record.get("path") if isinstance(record, dict) else None
        self.set_selection([path] if path else [])
        first = self.first_selected()
        if first != self.current_file:
            self.trace = None
            self.trace_error = ""
        self.current_file = first

    def edit_cell(self, record: Any, key: str, value: Any) -> None:
        """Write an edited table cell (the table's ``edited_call``): rating or notes.

        The table has already put *value* into the row; an invalid rating (not a whole
        number from 0 to :data:`RATING_MAX`) or an edit while a scan runs puts the stored
        value back.
        """
        if not isinstance(record, dict) or not record.get("path"):
            return
        path = record["path"]
        if self.busy:
            self._restore_cell(record, key)
            return
        if key == "rating":
            try:
                rating = int(value)
                if rating != float(value) or not 0 <= rating <= RATING_MAX:
                    raise ValueError(value)
            except (TypeError, ValueError):
                self._restore_cell(record, key)
                self.error_text = f"A rating is a whole number from 0 to {RATING_MAX}."
                return
            self.error_text = ""
            self.set_rating(path, rating)
        elif key == "notes":
            self.set_notes(path, str(value))

    def _restore_cell(self, record: dict, key: str) -> None:
        """Put the stored rating or notes of *record*'s file back into the row."""
        if key == "rating":
            record["rating"] = self.get_rating(record["path"])
        elif key == "notes":
            record["notes"] = self.get_notes(record["path"])

    def export_view(self) -> dict[str, Any]:
        """What the Browser page remembers: folder, include-subfolders, filter, bin window, y range, precompute."""
        return {
            "folder": str(self.current_folder) if self.current_folder else "",
            "include_subfolders": bool(self.include_subfolders),
            "rating_filter": self.rating_filter,
            "window_ms": float(self.window_ms),
            "y_min": float(self.y_min),
            "y_max": float(self.y_max),
            "precompute_after_scan": bool(self.precompute_after_scan),
        }

    def restore_view(self, state: dict | None) -> pathlib.Path | None:
        """Adopt :meth:`export_view` (invalid or missing entries keep the current value).

        Returns
        -------
        pathlib.Path or None
            The remembered folder if it still is one (the caller opens it), else ``None``.
        """
        state = state or {}
        if "include_subfolders" in state:
            self.include_subfolders = bool(state["include_subfolders"])
        if "precompute_after_scan" in state:
            self.precompute_after_scan = bool(state["precompute_after_scan"])
        if state.get("rating_filter") in FILTER_LABELS:
            self.rating_filter = state["rating_filter"]
        for key in ("window_ms", "y_min", "y_max"):
            try:
                if key in state:
                    setattr(self, key, float(state[key]))
            except (TypeError, ValueError):
                pass
        folder = str(state.get("folder") or "")
        if folder and pathlib.Path(folder).is_dir():
            return pathlib.Path(folder)
        return None

    # ── rating filter ─────────────────────────────────────────────────
    @property
    def rating_filter(self) -> str:
        """Label of the active rating filter (one of :data:`FILTER_LABELS`)."""
        idx = self.filter_index
        return FILTER_LABELS[idx] if 0 <= idx < len(FILTER_LABELS) else FILTER_LABELS[0]

    @rating_filter.setter
    def rating_filter(self, label: str) -> None:
        self.filter_index = FILTER_LABELS.index(label) if label in FILTER_LABELS else 0

    def filter_accept(self, rating: int) -> bool:
        """Return whether *rating* passes the rating filter."""
        idx = self.filter_index
        if idx == 0:
            return True
        elif idx == 1:
            return rating >= 1
        elif idx == 2:
            return rating >= 2
        elif idx == 3:
            return rating >= 3
        elif idx == 4:
            return rating == 0
        return True

    def accepts(
        self, rating: int, path: pathlib.Path, allowed_exts: set[str] | None = None
    ) -> bool:
        """Return whether a file passes the rating filter and the setup's extensions."""
        if allowed_exts is None:
            allowed_exts = self.allowed_exts()
        rating_ok = self.filter_accept(rating)
        ext_ok = (not allowed_exts) or (pathlib.Path(path).suffix.lower() in allowed_exts)
        return rating_ok and ext_ok

    def set_rating_filter(self, value: Any = None) -> None:
        """Apply the rating filter (``value`` is an index or a label, optional)."""
        if isinstance(value, str):
            self.rating_filter = value
        elif isinstance(value, int):
            self.filter_index = value
        self.apply_filter()

    def apply_filter(self) -> list[pathlib.Path]:
        """Re-filter :attr:`files` into :attr:`rows`; drop files that vanished.

        Returns
        -------
        list of pathlib.Path
            The files that no longer exist and were dropped from :attr:`files`.
        """
        missing: list[pathlib.Path] = []
        kept: list[dict[str, Any]] = []
        for row in self.files:
            p = pathlib.Path(row["path"])
            try:
                if not p.exists():
                    missing.append(p)
                    continue
            except Exception:
                pass  # be conservative and keep the row
            kept.append(row)
        self.files = kept
        allowed = self.allowed_exts()
        self.rows = [
            r for r in kept if self.accepts(int(r["rating"]), pathlib.Path(r["path"]), allowed)
        ]
        keep = {r["path"] for r in self.rows}
        self.selected_files = [s for s in self.selected_files if s in keep]
        # A file the filter hides (or that vanished) is no longer shown: the Qt browser cleared its
        # plot and annotation as soon as the selected row disappeared from the table.
        if self.current_file is not None and (
            not self.current_file.exists() or str(self.current_file) not in keep
        ):
            self.current_file = None
            self.trace = None
            self.trace_error = ""
        if self.current_folder is not None:
            self.status_text = (
                f"{len(self.rows)} of {len(self.files)} file(s) in {self.current_folder}"
            )
        self.notify("rows")
        return missing

    # ── folder, scan ──────────────────────────────────────────────────
    def open_folder(self, folder: str | pathlib.Path) -> bool:
        """Open *folder*: load its metadata and scan it.

        Returns
        -------
        bool
            ``False`` (and :attr:`error_text` set) when *folder* is not a directory
            or scanning failed.
        """
        folder = pathlib.Path(folder)
        try:
            if not folder.exists() or not folder.is_dir():
                self.error_text = f"Not a folder: {folder}"
                logger.warning("TraceBrowser: Selected path is not a folder: %s", folder)
                return False
            self.error_text = ""
            self.current_folder = folder
            logger.info("TraceBrowser: Opened folder %s", folder)
            self.load_meta()
            self.scan()
            return True
        except Exception as e:
            self.error_text = f"Failed to open {folder}: {e}"
            logger.exception("TraceBrowser: Failed to open folder %s: %s", folder, e)
            return False

    def set_include_subfolders(self, value: Any = None) -> None:
        """Set the include-subfolders flag (when given) and rescan the open folder."""
        if value is not None:
            self.include_subfolders = bool(value)
        if self.current_folder:
            self.request("scan")

    def scan(self) -> list[dict[str, Any]]:
        """Scan :attr:`current_folder` and fill :attr:`files` and :attr:`rows`.

        Lists the files of the setup's file types (optionally recursively),
        skipping anything under a ``.trash`` folder and image-like (CLSM) TTTR
        files, auto-selects the channels of the first usable file when no setup
        defined them, and reads each file's rating and notes from the metadata.

        Returns
        -------
        list of dict
            The visible :attr:`rows`.
        """
        if not self.current_folder:
            return []
        folder = self.current_folder
        files: list[pathlib.Path] = []
        allowed_exts = self.allowed_exts()
        recursive = bool(self.include_subfolders)
        try:
            it = folder.rglob("*") if recursive else folder.iterdir()
        except Exception:
            it = folder.iterdir()
        for p in sorted(it):
            try:
                if not p.is_file():
                    continue
                # Exclude anything under a hidden .trash within the current folder
                try:
                    relp = p.relative_to(folder)
                    if any(part == ".trash" for part in relp.parts):
                        continue
                except Exception:
                    pass
                if not allowed_exts or p.suffix.lower() in allowed_exts:
                    # Skip CLSM/image TTTR files; Trace Browser should only list non-image TTTRs
                    try:
                        if self.is_image_tttr(p):
                            continue
                    except Exception:
                        pass
                    files.append(p)
            except Exception:
                continue
        # Auto-initialize channels if setup was not used: detect from first TTTR
        try:
            if tttrlib is not None and self.selected_channels is None and files:
                for _p in files:
                    try:
                        if self.is_image_tttr(_p):
                            continue
                    except Exception:
                        pass
                    try:
                        tt = tttrlib.TTTR(str(_p))
                        chs = sorted(tt.get_used_routing_channels())
                        if chs:
                            self.selected_channels = chs
                            logger.info(
                                "TraceBrowser: Auto-selected channels from %s: %s", _p.name, chs
                            )
                            break
                    except Exception:
                        continue
        except Exception as _e:
            logger.debug("TraceBrowser: Auto channel initialization skipped: %s", _e)
        self.files = [self._make_row(p) for p in files]
        self.precompute["message"] = ""
        self.apply_filter()
        # The Qt browser selects the first row after every scan and plots it (or clears the
        # plot when nothing is listed), then precomputes the traces of the listed files.
        if self.rows:
            self.select_row(self.rows[0])
            self.precompute_pending = bool(self.precompute_after_scan)
        else:
            self.set_selection([])
            self.current_file = None
            self.trace = None
            self.trace_error = ""
        logger.debug(
            "TraceBrowser: Found %d files (%s), displaying %d after filter",
            len(files),
            "recursive" if recursive else "flat",
            len(self.rows),
        )
        self.status_text = f"{len(self.rows)} of {len(files)} file(s) in {folder}"
        self.notify("files")
        return self.rows

    def _make_row(self, p: pathlib.Path) -> dict[str, Any]:
        """Build the table row of file *p* from the disk and the metadata."""
        try:
            rel_txt = str(p.relative_to(self.current_folder))
        except Exception:
            rel_txt = p.name
        try:
            sz = p.stat().st_size
        except Exception:
            sz = 0
        rec = self.meta_get(p)
        return {
            "name": rel_txt,
            "path": str(p),
            "size": sz,
            "size_text": _trace.human_size(sz),
            "size_mb": float(sz) / (1024.0 * 1024.0),
            "rating": int(rec.get("rating", 0)),
            "notes": str(rec.get("annotation", "")),
        }

    def clear(self) -> None:
        """Clear the file list, the selection and the shown trace (files and metadata stay)."""
        self.files = []
        self.rows = []
        self.selected_files = []
        self.current_file = None
        self.trace = None
        self.trace_error = ""
        logger.info("TraceBrowser: Cleared file list")
        self.notify("rows")

    # ── image detection ───────────────────────────────────────────────
    def is_clsm_compatible(self, tttr_obj: Any) -> bool:
        """Probe whether *tttr_obj* supports CLSM imaging (an image-like dataset)."""
        try:
            clsm = tttrlib.CLSMImage(tttr_data=tttr_obj)
            intensity = getattr(clsm, "intensity", None)
            if intensity is None:
                return False
            # tttrlib builds a CLSMImage from *any* readable TTTR without raising; a file that
            # is not a scan gives an empty array (1, 0, 0). Only a non-empty pixel stack is an
            # image: testing ``is not None`` classed every single-molecule file as one and the
            # browser then listed no TTTR file at all.
            return int(np.asarray(intensity).size) > 0
        except Exception:
            return False

    def probe_image(self, path: pathlib.Path) -> bool:
        """The default image probe: open *path* with ``tttrlib`` and test CLSM support."""
        v = self._is_image_cache.get(path)
        if v is not None:
            return v
        if tttrlib is None:
            self._is_image_cache[path] = False
            return False
        try:
            tt = tttrlib.TTTR(str(path))
            v = self.is_clsm_compatible(tt)
            self._is_image_cache[path] = bool(v)
            return bool(v)
        except Exception:
            # On failure to open, treat as non-image to let other filters decide
            self._is_image_cache[path] = False
            return False

    def is_image_tttr(self, path: pathlib.Path) -> bool:
        """Return whether *path* is an image (CLSM) TTTR file (uses :attr:`image_probe` if set)."""
        if self.image_probe is not None:
            return bool(self.image_probe(path))
        return self.probe_image(path)

    # ── metadata: ratings and notes ───────────────────────────────────
    def load_meta(self) -> dict[str, dict]:
        """(Re)load the metadata of :attr:`current_folder` (RPC first, then the file)."""
        folder = self.current_folder
        if folder is None:
            self.meta = {}
            return self.meta
        self.meta = self.client.get_metadata(str(folder)) or _metadata.load_meta(folder)
        return self.meta

    def rel_key(self, path: pathlib.Path) -> str:
        """Return the metadata key of *path*: relative to the folder, else the file name."""
        try:
            if self.current_folder is not None:
                return str(pathlib.Path(path).resolve().relative_to(self.current_folder.resolve()))
        except Exception:
            pass
        try:
            return pathlib.Path(path).name
        except Exception:
            return str(path)

    def meta_get(self, path: pathlib.Path) -> dict:
        """Return the metadata record of *path* (an empty dict when there is none)."""
        key = self.rel_key(path)
        rec = self.meta.get(key)
        if rec is not None:
            return rec
        # Backward compatibility: older meta by filename only
        try:
            return self.meta.get(pathlib.Path(path).name, {})
        except Exception:
            return {}

    def meta_set(self, path: pathlib.Path, rec: dict) -> None:
        """Store *rec* as the metadata of *path* and push the metadata to the backend."""
        key = self.rel_key(path)
        self.meta[key] = rec
        if self.current_folder is not None:
            try:
                self.client.set_metadata(str(self.current_folder), dict(self.meta))
            except Exception:
                pass

    def flush_meta(self) -> bool:
        """Write the metadata to the folder's ``.trace_browser_meta.json``.

        Returns
        -------
        bool
            ``True`` when the file was written.
        """
        try:
            if self.current_folder:
                _metadata.save_meta(self.current_folder, self.meta)
                return True
        except Exception:
            pass
        return False

    def get_rating(self, path: str | pathlib.Path) -> int:
        """Return the 0..3 rating of *path* (0 when unrated)."""
        return int(self.meta_get(pathlib.Path(path)).get("rating", 0))

    def get_notes(self, path: str | pathlib.Path) -> str:
        """Return the notes (annotation) of *path*."""
        return str(self.meta_get(pathlib.Path(path)).get("annotation", ""))

    def set_rating(self, path: str | pathlib.Path, rating: int, flush: bool = True) -> None:
        """Set the rating of *path*, persist it and refresh the filtered rows.

        Parameters
        ----------
        path : str or pathlib.Path
            The file.
        rating : int
            New rating (0..3).
        flush : bool
            Also write the metadata file now (a host that debounces writes passes ``False``).
        """
        if not self.current_folder:
            return
        path = pathlib.Path(path)
        rec = self.meta_get(path)
        rec["rating"] = int(rating)
        self.meta_set(path, rec)
        self._sync_row(path, rating=int(rating))
        if flush:
            self.flush_meta()
        self.apply_filter()
        self.notify("meta")

    def set_notes(self, path: str | pathlib.Path, text: str, flush: bool = True) -> None:
        """Set the notes (annotation) of *path* and persist them."""
        if not self.current_folder:
            return
        path = pathlib.Path(path)
        rec = self.meta_get(path)
        rec["annotation"] = str(text)
        self.meta_set(path, rec)
        self._sync_row(path, notes=str(text))
        if flush:
            self.flush_meta()
        self.notify("meta")

    def _sync_row(self, path: pathlib.Path, **fields: Any) -> None:
        """Update the row of *path* in :attr:`files` (rows share the same dicts)."""
        sp = str(path)
        for row in self.files:
            if row["path"] == sp:
                row.update(fields)

    # ── selection ─────────────────────────────────────────────────────
    def set_selection(self, paths: list[str | pathlib.Path]) -> None:
        """Select the rows of *paths* (only files that are listed)."""
        listed = {r["path"] for r in self.rows}
        self.selected_files = [str(p) for p in paths if str(p) in listed]
        self.notify("selection")

    def first_selected(self) -> pathlib.Path | None:
        """Return the first selected file, or ``None``."""
        return pathlib.Path(self.selected_files[0]) if self.selected_files else None

    # ── y range ───────────────────────────────────────────────────────
    @property
    def y_range(self) -> tuple[float, float]:
        """The y range ordered as ``(low, high)`` (the two fields may be entered swapped)."""
        lo, hi = float(self.y_min), float(self.y_max)
        return (hi, lo) if lo > hi else (lo, hi)

    # ── traces and caches ─────────────────────────────────────────────
    def _cache_folder(self) -> str | None:
        return str(self.current_folder) if self.current_folder is not None else None

    def trace_signature(self, path: pathlib.Path, window_ms: float | None = None) -> str:
        """Return the cache signature of the trace of *path* for the current setup."""
        window_ms = self.window_ms if window_ms is None else window_ms
        return _trace.trace_signature(
            pathlib.Path(path), window_ms, self.setup_settings, self.selected_channels
        )

    def load_trace_cache(
        self, path: pathlib.Path, window_ms: float | None = None
    ) -> tuple[np.ndarray, np.ndarray, list[str]] | None:
        """Return the on-disk cached trace of *path*, or ``None``."""
        window_ms = self.window_ms if window_ms is None else window_ms
        folder = self.current_folder or pathlib.Path(path).parent
        return _trace.load_cached(
            pathlib.Path(path), folder, window_ms, self.setup_settings, self.selected_channels
        )

    def compute_trace_cached(
        self, path: pathlib.Path, window_ms: float | None = None
    ) -> tuple[np.ndarray, np.ndarray, list[str]]:
        """Return ``(time_axis, counts, labels)`` of *path*: memory cache, disk cache, compute.

        The computation and the on-disk cache are :func:`..core.trace.load_trace`
        (the same code the RPC backend runs); the in-memory cache is the model's.
        """
        path = pathlib.Path(path)
        window_ms = self.window_ms if window_ms is None else window_ms
        key = (path, self.trace_signature(path, window_ms))
        hit = self._trace_mem_cache.get(key)
        if hit is not None:
            return hit
        data = _trace.load_trace(
            str(path),
            time_window_ms=window_ms,
            setup_settings=self.setup_settings,
            selected_channels=self.selected_channels,
            cache_folder=self._cache_folder(),
        )
        result = (
            np.asarray(data["time_axis"], dtype=float),
            np.asarray(data["counts"], dtype=float),
            [str(label) for label in data["labels"]],
        )
        self._trace_mem_cache[key] = result
        return result

    def load_trace(
        self, path: str | pathlib.Path, window_ms: float | None = None
    ) -> tuple[np.ndarray, np.ndarray, list[str]]:
        """Load the binned trace of *path* (RPC client first, then in-process compute).

        Returns
        -------
        tuple
            ``(time_axis, counts, labels)``; *counts* is ``(bins, series)``.
        """
        path = pathlib.Path(path)
        window_ms = self.window_ms if window_ms is None else window_ms
        trace = self.client.load_trace(
            str(path),
            time_window_ms=window_ms,
            setup_settings=self.setup_settings,
            selected_channels=self.selected_channels,
            cache_folder=self._cache_folder(),
        )
        if trace is not None:
            return (
                np.asarray(trace.get("time_axis", []), dtype=float),
                np.asarray(trace.get("counts", []), dtype=float),
                [str(label) for label in trace.get("labels", [])],
            )
        return self.compute_trace_cached(path, window_ms)

    def show_file(self, path: str | pathlib.Path) -> bool:
        """Load the trace of *path* into :attr:`trace` and make it the current file.

        Returns
        -------
        bool
            ``False`` (and :attr:`error_text` set) when the trace could not be loaded.
        """
        path = pathlib.Path(path)
        try:
            time_axis, counts, labels = self.load_trace(path)
        except Exception as e:
            self.error_text = f"Failed to load {path.name}: {e}"
            logger.exception("TraceBrowser: Failed to load %s: %s", path, e)
            return False
        self.error_text = ""
        self.current_file = path
        self.set_trace(path, time_axis, counts, labels)
        return True

    def set_trace(
        self,
        path: str | pathlib.Path,
        time_axis: np.ndarray,
        counts: np.ndarray,
        labels: list[str],
        window_ms: float | None = None,
    ) -> None:
        """Make ``(time_axis, counts, labels)`` the shown :attr:`trace` of *path*.

        A host that loads the trace on a worker hands the result over here.

        Parameters
        ----------
        path : str or pathlib.Path
            The file the trace belongs to.
        time_axis, counts, labels
            What :meth:`load_trace` returned.
        window_ms : float, optional
            The bin window the trace was loaded with (default: the current one).
        """
        self.trace_error = ""
        self.trace = {
            "path": str(path),
            "time_axis": time_axis,
            "counts": counts,
            "labels": labels,
            "time_window_ms": float(self.window_ms if window_ms is None else window_ms),
        }
        self.notify("trace")

    def fail_trace(self, path: str | pathlib.Path, error: Any) -> None:
        """Record that the trace of *path* could not be loaded (shown in the plot area)."""
        self.trace = None
        self.trace_error = f"Failed to load {pathlib.Path(path).name}: {error}"
        self.notify("trace")

    def set_window_ms(self, value: Any = None) -> None:
        """Set the bin window (when given) and reload the trace of the current file."""
        if value is not None:
            self.window_ms = float(value)
        if self.current_file is not None:
            self.show_file(self.current_file)

    def warm_trace_cache(self, paths: list[pathlib.Path]) -> list[pathlib.Path]:
        """Load already-cached traces of *paths* into memory; return those still to compute."""
        todo: list[pathlib.Path] = []
        for p in paths:
            try:
                loaded = self.load_trace_cache(p, self.window_ms)
                if loaded is not None:
                    self._trace_mem_cache[(p, self.trace_signature(p, self.window_ms))] = loaded
                else:
                    todo.append(p)
            except Exception:
                todo.append(p)
        return todo

    def precompute_all_traces(
        self, progress: Callable[[int, int, pathlib.Path | None], bool | None] | None = None
    ) -> int:
        """Compute and cache the trace of every listed file for snappy browsing.

        Parameters
        ----------
        progress : callable, optional
            ``progress(i, total, path)`` before each file (return ``False`` to stop)
            and ``progress(total, total, None)`` once at the end.

        Returns
        -------
        int
            Number of traces computed.
        """
        if self.current_folder is None or tttrlib is None:
            return 0
        paths: list[pathlib.Path] = []
        for row in self.rows:
            pth = pathlib.Path(row["path"])
            # Defensive: skip image files even if present in the rows for any reason
            try:
                if self.is_image_tttr(pth):
                    continue
            except Exception:
                pass
            paths.append(pth)
        if not paths:
            return 0
        todo = self.warm_trace_cache(paths)
        if not todo:
            return 0
        computed = 0
        try:
            for i, p in enumerate(todo):
                if progress is not None and progress(i, len(todo), p) is False:
                    break
                try:
                    self.compute_trace_cached(p, self.window_ms)
                    computed += 1
                except Exception:
                    pass
        finally:
            if progress is not None:
                progress(len(todo), len(todo), None)
        self.notify("traces")
        return computed

    def precompute_traces(self) -> None:
        """Ask the host to precompute the traces of the listed files now (the *Precompute* button)."""
        self.precompute_pending = True

    def stop_precompute(self) -> None:
        """Ask the running precompute job to stop after the file it is working on."""
        self.precompute["cancel"] = True

    def prepare_precompute(self) -> None:
        """Mark the precompute job as running (the host calls this before it starts the worker)."""
        self.precompute.update(running=True, cancel=False, done=0, total=0, name="", message="")

    def run_precompute(self) -> int:
        """Run :meth:`precompute_all_traces` with progress kept in :attr:`precompute`.

        Meant for a worker thread: the draw loop reads :attr:`precompute`; :meth:`stop_precompute`
        makes the loop stop after the current file.

        Returns
        -------
        int
            Number of traces computed.
        """
        state = self.precompute
        state.update(running=True, done=0, total=0, name="")
        computed = 0

        def progress(i: int, total: int, path: pathlib.Path | None) -> bool:
            state["total"] = total
            if path is not None:
                state.update(done=i, name=path.name)
            return not state["cancel"]

        try:
            computed = self.precompute_all_traces(progress)
        finally:
            if state["cancel"]:
                state["message"] = f"Precompute stopped after {computed} trace(s)."
            elif state["total"] == 0:
                state["message"] = "No traces to compute (all cached)."
            else:
                state["message"] = f"Precomputed {computed} trace(s)."
            state["running"] = False
        return computed

    def clear_caches(self) -> int:
        """Clear the in-memory caches and the folder's on-disk trace cache.

        Returns
        -------
        int
            Number of cache directories removed (0 or 1).
        """
        self._trace_mem_cache.clear()
        self._is_image_cache.clear()
        removed_dirs = 0
        base = self.current_folder
        if base and isinstance(base, pathlib.Path):
            cache_dir = base / CACHE_DIRNAME
            try:
                if cache_dir.exists() and cache_dir.is_dir():
                    shutil.rmtree(str(cache_dir), ignore_errors=True)
                    removed_dirs += 1
            except Exception:
                pass
        logger.info("TraceBrowser: Cleared caches (memory + %d dir(s) removed)", removed_dirs)
        self.status_text = "Trace caches cleared."
        self.notify("caches")
        return removed_dirs

    # ── selection set, export, delete and hand-offs (the Qt toolbar's actions) ──────────
    def selection_paths(self) -> list[pathlib.Path]:
        """The files the export, delete and hand-off actions work on, in table order.

        The table's multi-row selection (:attr:`multi_selection`) when it holds listed files,
        else the single selected row (:attr:`selected_files`).
        """
        listed = [r["path"] for r in self.rows]
        known = set(listed)
        multi = {p for p in self.multi_selection if p in known}
        if multi:
            return [pathlib.Path(p) for p in listed if p in multi]
        return [pathlib.Path(p) for p in self.selected_files if p in known]

    def select_all_files(self) -> None:
        """Select every listed file (the table highlights them; the actions use all of them)."""
        self.multi_selection = [r["path"] for r in self.rows]
        self.select_all_requested = True
        self.status_text = f"{len(self.multi_selection)} file(s) selected."

    def _need_selection(self, what: str) -> list[pathlib.Path] | None:
        """The selected files, or ``None`` with a hint in :attr:`status_text` when there are none."""
        paths = self.selection_paths()
        if not paths:
            self.status_text = f"Select one or more files first, then {what}."
            return None
        return paths

    @property
    def docx_available(self) -> bool:
        """Whether the optional ``python-docx`` package can be imported."""
        return _docx_modules() is not None

    @staticmethod
    def ndx_available() -> bool:
        """Whether the ndX components can be imported (the Qt tool tested the same)."""
        return importlib.util.find_spec("ndxplorer") is not None

    # -- Export (copy the files) -----------------------------------------------------
    def export_selected(self) -> None:
        """Ask the host for a destination folder to copy the selected files to (*Export*)."""
        if self._need_selection("press Export") is not None:
            self.dialog = "export"

    def export_files(self, dest: str | pathlib.Path, paths: list | None = None) -> int:
        """Copy *paths* (default: the selection) into the folder *dest* with their metadata kept.

        The Qt tool's ``_on_export`` (``shutil.copy2`` into the chosen folder, flat, by file
        name) for the selected files instead of every listed file. A name that two selected
        files share (files of different sub-folders) gets ``__2``, ``__3`` ... so the second
        does not overwrite the first; a file that already exists in *dest* from before is
        overwritten, as in Qt.

        Returns
        -------
        int
            Number of files copied.
        """
        files = [pathlib.Path(p) for p in (paths if paths is not None else self.selection_paths())]
        out = pathlib.Path(dest)
        if not files:
            self.status_text = "Nothing selected to export."
            return 0
        copied = 0
        used: set[str] = set()
        try:
            out.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            self.error_text = f"Cannot create {out}: {exc}"
            return 0
        for src in files:
            name = src.name
            target = out / name
            n = 2
            while name in used:
                name = f"{src.stem}__{n}{src.suffix}"
                target = out / name
                n += 1
            used.add(name)
            try:
                if target.exists() and src.resolve() == target.resolve():
                    continue  # exporting a file onto itself
                shutil.copy2(str(src), str(target))
                copied += 1
            except Exception as exc:
                logger.warning("TraceBrowser: Failed to copy %s to %s: %s", src, out, exc)
        self.error_text = ""
        self.status_text = f"Exported {copied}/{len(files)} file(s) to {out}"
        if copied:
            self.actions_done += 1
        self.notify("export")
        return copied

    # -- CSV ---------------------------------------------------------------------------
    def export_csv(self) -> None:
        """Ask the host for a destination folder for the CSV traces of the selection (*CSV*)."""
        if self._need_selection("press CSV") is not None:
            self.dialog = "csv"

    def export_csv_files(self, dest: str | pathlib.Path, paths: list | None = None) -> list[str]:
        """Write the binned trace of each of *paths* (default: the selection) as CSV into *dest*.

        As the Qt ``_on_export_csv``: the plugin's RPC export first (one
        ``<stem>_trace.csv`` per file: the header ``time_ms`` then one column per series, one
        row per bin at the current bin window); when that gives nothing, the local fallback
        (``<stem>_bin<window>ms.csv``, header ``time_s``, image-like files skipped).

        Returns
        -------
        list of str
            The CSV files written.
        """
        files = [pathlib.Path(p) for p in (paths if paths is not None else self.selection_paths())]
        out = pathlib.Path(dest)
        if not files:
            self.status_text = "Nothing selected to export."
            return []
        out.mkdir(parents=True, exist_ok=True)
        window_ms = float(self.window_ms)
        written: list[str] = []
        try:
            written = list(
                self.client.export_csv(
                    [str(p) for p in files],
                    str(out),
                    time_window_ms=window_ms,
                    setup_settings=self.setup_settings,
                    selected_channels=self.selected_channels,
                )
                or []
            )
        except Exception:
            logger.debug("TraceBrowser: RPC CSV export failed; falling back to local export")
        if not written:
            written = self._export_csv_local(files, out, window_ms)
        self.error_text = ""
        self.status_text = f"Exported {len(written)}/{len(files)} CSV file(s) to {out}"
        if written:
            self.actions_done += 1
        self.notify("export")
        return written

    def _export_csv_local(
        self, files: list[pathlib.Path], out: pathlib.Path, window_ms: float
    ) -> list[str]:
        """The Qt tool's local CSV fallback (cut and pasted from ``_on_export_csv``)."""
        written: list[str] = []
        for p in files:
            try:
                if self.is_image_tttr(p):
                    continue
                time_axis, padded, labels = self.compute_trace_cached(p, window_ms)
                if time_axis is None or padded is None or len(time_axis) == 0:
                    continue
                safe_labels = [
                    str(l).replace("\n", " ").replace("\r", " ").replace(",", ";") for l in labels
                ]
                header = ["time_s"] + safe_labels
                bin_tag = (f"{window_ms:g}").replace(".", "p")
                csv_path = out / f"{p.stem}_bin{bin_tag}ms.csv"
                with open(csv_path, "w", newline="", encoding="utf-8") as f:
                    writer = csv.writer(f)
                    writer.writerow(header)
                    nb = int(padded.shape[0])
                    if len(padded.shape) == 1:
                        for i in range(nb):
                            writer.writerow([float(time_axis[i]), float(padded[i])])
                    else:
                        nc = int(padded.shape[1])
                        for i in range(nb):
                            writer.writerow(
                                [float(time_axis[i])] + [float(padded[i, j]) for j in range(nc)]
                            )
                written.append(str(csv_path))
            except Exception as exc:
                logger.warning("TraceBrowser: Failed to export CSV for %s: %s", p, exc)
        return written

    # -- DOCX report ---------------------------------------------------------------------
    def export_docx(self) -> None:
        """Write the DOCX report of the selected files (*DOCX*; needs ``python-docx``)."""
        paths = self._need_selection("press DOCX")
        if paths is None:
            return
        if not self.docx_available:
            self.error_text = (
                "python-docx is not installed. Please install 'python-docx' to enable DOCX export."
            )
            return
        self.request("write_docx", [str(p) for p in paths])

    def write_docx(self, paths: list | None = None) -> pathlib.Path | None:
        """Write ``<folder name>.docx`` into the open folder: one section per file.

        The Qt tool's ``_on_export_docx``: a heading, then per file its name, folder, rating,
        annotation and a picture of its trace (drawn here from the loaded trace with
        :mod:`emtk.figure`, which needs no toolkit; no picture when the trace cannot be read).

        Returns
        -------
        pathlib.Path or None
            The saved file, or ``None`` (and :attr:`error_text` set) when it could not be written.
        """
        modules = _docx_modules()
        if modules is None:
            self.error_text = (
                "python-docx is not installed. Please install 'python-docx' to enable DOCX export."
            )
            return None
        Document, Inches = modules
        files = [pathlib.Path(p) for p in (paths if paths is not None else self.selection_paths())]
        if not files:
            self.status_text = "Nothing selected to export."
            return None
        folder = self.current_folder or files[0].parent
        save_path = folder / ((folder.name or "traces") + ".docx")
        tmpdir = pathlib.Path(tempfile.mkdtemp(prefix="trace_export_"))
        try:
            doc = Document()
            doc.add_heading("Trace Browser Export", level=1)
            for p in files:
                rec = self.meta_get(p)
                doc.add_heading(p.name, level=2)
                doc.add_paragraph(f"Folder: {p.parent}")
                doc.add_paragraph(f"Rating: {int(rec.get('rating', 0))}")
                annotation = rec.get("annotation", "")
                if annotation:
                    doc.add_paragraph(f"Annotation: {annotation}")
                image = self._render_trace_png(p, tmpdir / f"{p.stem}.png")
                if image is not None:
                    doc.add_picture(str(image), width=Inches(6))
                doc.add_paragraph("")
            doc.save(str(save_path))
        except Exception as exc:
            logger.exception("TraceBrowser: Failed to save DOCX %s: %s", save_path, exc)
            self.error_text = f"Failed to save DOCX: {exc}"
            return None
        finally:
            shutil.rmtree(str(tmpdir), ignore_errors=True)
        self.error_text = ""
        self.status_text = f"Saved: {save_path}"
        self.actions_done += 1
        self.notify("export")
        return save_path

    def _render_trace_png(self, path: pathlib.Path, png: pathlib.Path) -> pathlib.Path | None:
        """Draw the trace of *path* into *png* (:mod:`emtk.figure`); ``None`` when that fails."""
        try:
            from emtk.figure import Figure

            time_axis, counts, labels = self.load_trace(path, self.window_ms)
            counts = np.asarray(counts, dtype=float)
            if counts.ndim == 1:
                counts = counts[:, None]
            ax = Figure(size=(700, 300)).ax()
            for j, label in enumerate(labels[: counts.shape[1]]):
                ax.line(time_axis, counts[:, j], width=0.8, label=str(label))
            ax.set_labels(x="Time (s)", y=f"Counts / {self.window_ms:g} ms")
            ax.set_title(path.name)
            if labels:
                ax.legend("ne")
            ax.figure.save(png)
            return png
        except Exception as exc:
            logger.debug("TraceBrowser: No trace picture for %s: %s", path, exc)
            return None

    # -- Delete (move to .trash) --------------------------------------------------------
    def plan_delete(self, paths: list | None = None) -> list[pathlib.Path]:
        """The files a delete moves: *paths* plus every sibling sharing a file name stem.

        As the Qt ``_on_delete_selected`` (a measurement's ``.spc`` and its companion files
        move together). Only existing files inside the open folder and outside ``.trash`` are
        taken; nothing else is ever touched.
        """
        folder = self.current_folder
        if folder is None:
            return []
        base = folder.resolve()
        chosen: dict[pathlib.Path, None] = {}

        def add(fp: pathlib.Path) -> None:
            try:
                if not fp.is_file():
                    return
                rel = fp.resolve().relative_to(base)
            except Exception:
                return  # outside the open folder: never touched
            if any(part == ".trash" for part in rel.parts):
                return
            chosen[fp] = None

        for item in paths if paths is not None else self.selection_paths():
            p = pathlib.Path(item)
            add(p)
            try:
                for sib in p.parent.iterdir():
                    if sib.name.startswith(p.stem + ".") and sib.stem == p.stem:
                        add(sib)
            except Exception:
                pass
        return sorted(chosen)

    def delete_selected(self) -> None:
        """Ask for confirmation to move the selected files to ``.trash`` (*Delete selected*)."""
        paths = self._need_selection("press Delete selected")
        if paths is None:
            return
        move = self.plan_delete(paths)
        if not move:
            self.status_text = "Nothing to move: the selected files are not in the open folder."
            return
        names = [p.name for p in move]
        shown = ", ".join(names[:6]) + (f", ... ({len(names) - 6} more)" if len(names) > 6 else "")
        self.confirm = {
            "kind": "delete",
            "title": "Move to .trash?",
            "message": (
                f"Move {len(move)} file(s) to the .trash folder of {self.current_folder}? "
                f"The selection plus the files that share its name: {shown}. "
                "Nothing is deleted for good: move a file back by hand to list it again."
            ),
            "yes": "Move to .trash",
            "paths": [str(p) for p in move],
        }

    def delete_from_table(self, record: Any = None) -> None:
        """The table's Delete key (once per selected row): ask to move the selection to ``.trash``."""
        if self.confirm is None and not self.busy:
            self.delete_selected()

    def confirm_yes(self) -> None:
        """Carry out the pending confirmation."""
        pending, self.confirm = self.confirm, None
        if pending and pending.get("kind") == "delete":
            self.request("delete_files", list(pending["paths"]))

    def confirm_no(self) -> None:
        """Drop the pending confirmation; nothing changes."""
        if self.confirm is not None:
            self.confirm = None
            self.status_text = "Cancelled: no file was moved."

    def trash_dir(self) -> pathlib.Path | None:
        """The ``.trash`` folder of the open folder (created on use)."""
        base = self.current_folder
        if base is None:
            return None
        trash = base / ".trash"
        try:
            trash.mkdir(parents=True, exist_ok=True)
        except Exception:
            pass
        return trash

    def delete_files(self, paths: list) -> int:
        """Move *paths* into ``.trash`` (keeping sub-folders), drop their metadata, refresh the list.

        The Qt ``_on_delete_selected`` body: an existing target gets a ``__<timestamp>``
        suffix instead of being overwritten; the selection moves to the row nearest the first
        removed one; the plot clears when the shown file moved.

        Returns
        -------
        int
            Number of files moved.
        """
        folder = self.current_folder
        trash = self.trash_dir()
        if folder is None or trash is None:
            return 0
        anchor = None
        moved_set = {str(p) for p in paths}
        for i, row in enumerate(self.rows):
            if row["path"] in moved_set:
                anchor = i
                break
        base = folder.resolve()
        moved = 0
        for p in (pathlib.Path(x) for x in sorted(moved_set)):
            try:
                if not p.is_file():
                    continue
                rel = p.resolve().relative_to(base)  # outside the folder: not moved
                if any(part == ".trash" for part in rel.parts):
                    continue
                dest = trash / rel
                dest.parent.mkdir(parents=True, exist_ok=True)
                final = dest
                if final.exists():
                    ts = time.strftime("%Y%m%d-%H%M%S")
                    final = final.with_name(f"{final.stem}__{ts}{final.suffix}")
                shutil.move(str(p), str(final))
                logger.info("TraceBrowser: moved to trash: '%s' -> '%s'", p, final)
                moved += 1
                self.meta.pop(self.rel_key(p), None)
            except Exception as exc:
                logger.warning("TraceBrowser: Failed to move %s to .trash: %s", p, exc)
        self.flush_meta()
        self.multi_selection = []
        self.apply_filter()
        if self.rows and self.current_file is None and anchor is not None:
            self.select_row(self.rows[min(anchor, len(self.rows) - 1)])
        self.error_text = ""
        self.status_text = f"Moved {moved} file(s) to {trash}"
        if moved:
            self.actions_done += 1
        self.notify("rows")
        return moved

    # -- hand-offs to other tools ------------------------------------------------------
    def hand_off(self, name: str, payload: dict[str, Any]) -> None:
        """Record a request for the host (name and payload) in :attr:`requests`.

        The app passes new requests to its host callback (the ChiSurf main window) from the
        draw thread; the model never opens another window itself.
        """
        self.requests.append({"name": name, "payload": payload})
        self.actions_done += 1
        self.notify("request")

    def _channels_for(self, path: pathlib.Path) -> list[int]:
        """The routing channels of a hand-off: the setup's, else those used in *path* (Qt: ``[0, 2]``)."""
        if self.selected_channels is not None:
            return list(self.selected_channels)
        try:
            return sorted(int(c) for c in tttrlib.TTTR(str(path)).get_used_routing_channels())
        except Exception as exc:
            logger.warning("TraceBrowser: Failed to detect channels, using default [0, 2]: %s", exc)
            return [0, 2]

    def open_intensity_trace(self) -> None:
        """Hand the first selected trace to *Intensity Trace Analysis* (*HMM*)."""
        paths = self._need_selection("press HMM")
        if paths is None:
            return
        first = paths[0]
        self.hand_off(
            HANDOFF_INTENSITY_TRACE,
            {
                "file": str(first),
                "window_ms": float(self.window_ms),
                "setup_settings": self.setup_settings,
                "selected_channels": self._channels_for(first),
            },
        )
        self.status_text = f"Sent {first.name} to Intensity Trace Analysis."

    def open_time_window(self) -> None:
        """Hand the first selected file and the bin window to *TTTR Time Window* (*TW*)."""
        paths = self._need_selection("press TW")
        if paths is None:
            return
        first = paths[0]
        self.hand_off(
            HANDOFF_TIME_WINDOW, {"files": [str(first)], "window_ms": float(self.window_ms)}
        )
        self.status_text = f"Sent {first.name} to TTTR Time Window."

    def open_ndxplorer(self) -> None:
        """Write the burst analysis of the first selected file, then ask the host to open ndX (*NDX*)."""
        paths = self._need_selection("press NDX")
        if paths is None:
            return
        self.request("prepare_ndx", str(paths[0]))

    def prepare_ndx(self, source: str | pathlib.Path) -> pathlib.Path | None:
        """The Qt one-click ndX pipeline, up to opening the window.

        Cut from ``_on_open_in_ndxplorer``: window the file by the bin window, write the burst
        table (``bi4_bur/<stem>.bur``) and ``Info`` beside the data in ``<stem>_TW_<ms>ms`` and
        record the hand-off ``open_ndxplorer`` with that folder. (The Qt tool also saved the
        windows to a temporary ``.bst`` file that nothing read; that is not written here.)

        Returns
        -------
        pathlib.Path or None
            The analysis folder, or ``None`` (and :attr:`error_text` set) on failure.
        """
        src = pathlib.Path(source)
        if tttrlib is None:
            self.error_text = "tttrlib is not available."
            return None
        if not self.ndx_available():
            self.error_text = "ndX components are not available."
            return None
        tw_ms = float(self.window_ms)
        tw_s = tw_ms / 1000.0
        try:
            from chisurf.core.fio.fluorescence import burst as burstio
        except Exception:
            burstio = None
        try:
            from chisurf.plugins.tttr.tttr_time_windows.api.selection import compute_bids_from_tttr
        except Exception:
            compute_bids_from_tttr = None
        try:
            tttr = tttrlib.TTTR(str(src))
            if compute_bids_from_tttr is not None:
                bids = compute_bids_from_tttr(tttr, tw_s)
            else:
                mt = tttr.macro_times
                res = float(getattr(tttr.header, "macro_time_resolution", 0.0)) or float(
                    getattr(tttr, "macro_time_resolution", 0.0)
                )
                if res <= 0:
                    raise RuntimeError("Macro time resolution unavailable from TTTR header")
                clocks_per_bin = max(1, int(np.floor(tw_s / res)))
                max_clock = int(mt.max()) if len(mt) else 0
                edges = np.arange(0, max_clock + 1, clocks_per_bin, dtype=np.int64)
                starts = np.searchsorted(mt, edges, side="left")
                stops = np.searchsorted(mt, edges + clocks_per_bin, side="left")
                bids = np.stack([starts, stops], axis=1)
            if bids is None or getattr(bids, "size", 0) == 0:
                self.error_text = "No data to compute burst IDs."
                return None
            analysis_dir = src.parent / f"{src.stem}_TW_{tw_ms:.0f}ms"
            bi4_bur_dir = analysis_dir / "bi4_bur"
            (analysis_dir / "Info").mkdir(parents=True, exist_ok=True)
            bi4_bur_dir.mkdir(parents=True, exist_ok=True)
            mt_max = int(np.max(tttr.micro_times)) + 1 if len(tttr) > 0 else 0
            full = (0, mt_max if mt_max > 0 else 4096)
            settings = self.setup_settings
            if isinstance(settings, dict) and settings.get("detectors"):
                detectors = settings["detectors"]
            else:
                try:
                    chs = sorted(tttr.get_used_routing_channels())
                except Exception:
                    chs = []
                detectors = {"all": {"chs": chs, "micro_time_ranges": [full]}}
            windows = {"all": full}
            start_stop = [(int(a), int(b)) for a, b in np.asarray(bids).tolist()]
            if burstio is None:
                raise ImportError("burst utilities unavailable")
            df = burstio.generate_burst_dataframe(
                start_stop, str(src.name), tttr, windows, detectors
            )
            burstio.write_dataframe_to_bur(df, str(bi4_bur_dir / f"{src.stem}.bur"))
            try:
                max_macro_time = 0.0
                res = float(getattr(tttr.header, "macro_time_resolution", 0.0)) or float(
                    getattr(tttr, "macro_time_resolution", 0.0)
                )
                if len(tttr) > 0 and res > 0:
                    max_macro_time = float(tttr.macro_times.max()) * res
                burstio.write_mti_summary(
                    src, analysis_dir, max_macro_time=max_macro_time, append=True
                )
            except Exception as exc:
                logger.debug("TraceBrowser: MTI write skipped: %s", exc)
        except Exception as exc:
            logger.exception("TraceBrowser: One-click NDX workflow failed: %s", exc)
            self.error_text = f"One-click workflow failed: {exc}"
            return None
        self.error_text = ""
        self.hand_off(
            HANDOFF_NDXPLORER,
            {"folder": str(analysis_dir), "file": str(src), "window_ms": tw_ms},
        )
        self.status_text = f"Wrote {analysis_dir.name}; asked the host to open ndX."
        return analysis_dir
