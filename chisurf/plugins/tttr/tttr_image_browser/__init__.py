"""TTTR Image Browser Plugin Workspace Widget.

Browse TTTR files in a folder and preview intensity images for all
DetectorWizard-defined detector windows.  The browse experience (file list,
ratings, annotations, mosaic canvas, per-tile labels) is the reusable AutoForm
``image_browser`` section over a Qt-free :class:`ImageBrowserViewModel`; this
widget is the standalone shell (detector-setup page + toolbar-facing actions).
"""

from __future__ import annotations

import logging
import pathlib
import tempfile

import numpy as np
from qtpy.QtWidgets import (
    QFileDialog,
    QLabel,
    QMessageBox,
    QPushButton,
    QVBoxLayout,
    QWidget,
)

from chisurf.gui.autoform import AutoForm
from chisurf.gui.autoform.sections.image_browser_section import ImageBrowserWidget
from chisurf.gui.widgets.wizard.tttr_channeldefinition import DetectorWizardPage

from .gui.view_model import ImageBrowserViewModel

try:
    from docx import Document
    from docx.shared import Inches
except Exception:
    Document = None
    Inches = None

try:
    from chisurf.gui.misc_helpers import persist_plugin_state
except ImportError:
    persist_plugin_state = lambda n: lambda c: c  # noqa: E731

_log = logging.getLogger(__name__)

name = "Imaging:Tools:Image Browser"
META_FILENAME = ".image_browser_meta.json"
CACHE_DIR_NAME = ".tttr_image_cache"


@persist_plugin_state("tttr_image_browser")
class TTTRImageBrowser(QWidget):
    """Workspace widget for the TTTR Image Browser (detector setup + AutoForm browser)."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("TTTR Image Browser")
        self.model = ImageBrowserViewModel()

        root = QVBoxLayout(self)

        # ── Page 0: detector setup (unchanged wizard) ──
        self.page0 = QWidget(self)
        p0 = QVBoxLayout(self.page0)
        p0.addWidget(QLabel("Setup definition (DetectorWizard)", self.page0))
        self.detector_page = DetectorWizardPage(
            show_help=False,
            show_setups_file=True,
            show_setup_selection=True,
            show_tttr_reading=True,
            show_tables=True,
            show_add_inputs=True,
        )
        p0.addWidget(self.detector_page)
        self.btn_continue = QPushButton("Use setup and continue →", self.page0)
        self.btn_continue.clicked.connect(self._on_continue)
        p0.addWidget(self.btn_continue)

        # ── Page 1: AutoForm image browser ──
        self.page1 = QWidget(self)
        p1 = QVBoxLayout(self.page1)
        p1.setContentsMargins(0, 0, 0, 0)
        self.auto_form = AutoForm(self.model)
        p1.addWidget(self.auto_form)

        root.addWidget(self.page0)
        root.addWidget(self.page1)
        self.page1.hide()

        self.model.add_observer(self._on_model_event)
        self.setAcceptDrops(True)

    # ── browser widget access ──
    @property
    def _browser(self) -> ImageBrowserWidget | None:
        return self.auto_form.findChild(ImageBrowserWidget)

    def _on_model_event(self, event: str) -> None:
        for fn in (self.auto_form.sync_fields, self.auto_form.refresh_plots):
            try:
                fn()
            except Exception:
                _log.debug("image-browser refresh failed", exc_info=True)

    # ── compat surface used by the toolbar / coordinator ──
    @property
    def _current_file(self):
        return pathlib.Path(self.model.current_file) if self.model.current_file else None

    def _selected_paths(self) -> list[pathlib.Path]:
        return [pathlib.Path(p) for p in self.model.selected_files]

    # ── setup page ──
    def _on_continue(self):
        self.model.setup_settings = self.detector_page.get_settings()
        self.page0.hide()
        self.page1.show()

    def _on_back_to_setup(self):
        self.page1.hide()
        self.page0.show()

    def apply_setup_settings(self, payload: dict) -> None:
        """Apply a shared detector definition and skip the local setup page."""
        if not payload:
            return
        try:
            self.detector_page.load_data_into_tables(payload)
        except Exception:  # pragma: no cover - best-effort GUI sync
            _log.debug("Could not load shared setup into detector page", exc_info=True)
        self.model.apply_setup_settings(payload)
        self.page0.hide()
        self.page1.show()

    # ── toolbar actions (drive the view-model) ──
    def _on_pick_folder(self):
        folder = QFileDialog.getExistingDirectory(self, "Select folder with TTTR files")
        if folder:
            self.model.open_folder(folder)

    def _on_clear(self):
        self.model.clear()

    def _on_subfolders_toggled(self, checked: bool = None):
        self.model.set_recursive(bool(checked) if checked is not None else not self.model.recursive)

    def _on_clear_caches(self):
        self.model.clear_caches()
        base = pathlib.Path(self.model.current_folder) if self.model.current_folder else None
        if base and base.exists():
            import shutil

            for d in base.rglob(CACHE_DIR_NAME):
                shutil.rmtree(str(d), ignore_errors=True)
            QMessageBox.information(self, "Caches cleared", "Image caches have been cleared.")

    def _on_export(self):
        paths = self._selected_paths()
        if not paths:
            return
        out_dir = QFileDialog.getExistingDirectory(self, "Select destination folder")
        if not out_dir:
            return
        out = pathlib.Path(out_dir)
        out.mkdir(parents=True, exist_ok=True)
        from shutil import copy2

        for p in paths:
            try:
                copy2(str(p), str(out / p.name))
            except Exception as exc:
                _log.warning("Failed to copy %s: %s", p, exc)

    def _on_save_tiff(self):
        paths = self._selected_paths()
        if not paths:
            return
        out_dir = QFileDialog.getExistingDirectory(
            self, "Select destination folder for TIFF stacks"
        )
        if not out_dir:
            return
        try:
            self.model._client.export_tiff(
                [str(p) for p in paths], out_dir, self.model.setup_settings
            )
        except Exception as exc:
            _log.warning("TIFF export failed: %s", exc)

    def _on_export_docx(self):
        entries = self.model.file_entries()
        folder = self.model.current_folder
        if not entries or Document is None or not folder:
            return
        folder = pathlib.Path(folder)
        docx_path = folder / f"{folder.name or 'images'}.docx"
        doc = Document()
        doc.add_heading("TTTR Image Browser Export", level=1)
        temp_dir = pathlib.Path(tempfile.mkdtemp(prefix="img_docx_"))
        from qtpy.QtGui import QImage

        from chisurf.plugins.tttr.tttr_image_browser.core.image import get_magma_lut

        try:
            for entry in entries:
                path = entry["id"]
                doc.add_heading(pathlib.Path(path).name, level=2)
                doc.add_paragraph(f"Rating: {self.model.rating_of(path)}")
                doc.add_paragraph(f"Annotation: {self.model.note_of(path)}")
                self.model.current_file = path
                mosaic = self.model.current_image()
                if mosaic is None:
                    continue
                h, w = mosaic.shape
                lut = get_magma_lut()
                if lut is not None:
                    rgb = np.ascontiguousarray(lut[mosaic])
                    qimg = QImage(rgb.data, w, h, 3 * w, QImage.Format_RGB888)
                else:
                    qimg = QImage(mosaic.data, w, h, w, QImage.Format_Grayscale8)
                img_path = temp_dir / f"{pathlib.Path(path).stem}.png"
                if qimg.save(str(img_path), "PNG"):
                    doc.add_picture(str(img_path), width=Inches(6) if Inches else None)
            doc.save(str(docx_path))
            QMessageBox.information(self, "DOCX Export", f"Saved DOCX to {docx_path}")
        except Exception as exc:
            _log.exception("DOCX export failed: %s", exc)
        finally:
            import shutil

            shutil.rmtree(str(temp_dir), ignore_errors=True)

    # ── drag-drop a folder onto the workspace ──
    def _first_dropped_directory(self, event):
        md = event.mimeData()
        if md and md.hasUrls():
            for url in md.urls():
                local = url.toLocalFile()
                if local and pathlib.Path(local).is_dir():
                    return pathlib.Path(local)
        return None

    def dragEnterEvent(self, event):  # noqa: N802
        """Accept a drag that carries a folder URL."""
        (event.acceptProposedAction() if self._first_dropped_directory(event) else event.ignore())

    def dropEvent(self, event):  # noqa: N802
        """Open the first dropped directory."""
        folder = self._first_dropped_directory(event)
        if folder is not None:
            event.acceptProposedAction()
            self.model.open_folder(str(folder))
        else:
            event.ignore()


if __name__ == "plugin":
    from chisurf.plugins.tttr.tttr_image_browser.gui.tool import TTTRImageBrowserTool

    window = TTTRImageBrowserTool()
    window.show()
    window.raise_()
    window.activateWindow()
